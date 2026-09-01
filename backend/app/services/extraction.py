from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook


@dataclass
class ExtractionResult:
    text: str
    status: str
    error: str | None = None
    metadata: dict = field(default_factory=dict)


class PDFExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        document = fitz.open(stream=content, filetype="pdf")
        parts = [page.get_text() for page in document]
        document.close()
        return ExtractionResult(text="\n".join(parts).strip(), status="extracted", metadata={"pages": len(parts)})


class DOCXExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from io import BytesIO

        document = Document(BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs if p.text)
        return ExtractionResult(text=text.strip(), status="extracted")


class XLSXExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from io import BytesIO

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append("\t".join(cells))
        workbook.close()
        return ExtractionResult(text="\n".join(lines).strip(), status="extracted")


class ImageExtractor:
    def extract(self, content: bytes, filename: str) -> ExtractionResult:
        return ExtractionResult(
            text="",
            status="extracted",
            metadata={"note": "Phase 2 stores images without OCR", "filename": filename, "bytes": len(content)},
        )


class DocumentExtractor:
    def extract(self, content: bytes, filename: str, content_type: str | None = None) -> ExtractionResult:
        suffix = Path(filename).suffix.lower()
        try:
            if suffix == ".pdf":
                return PDFExtractor().extract(content)
            if suffix == ".docx":
                return DOCXExtractor().extract(content)
            if suffix == ".xlsx":
                return XLSXExtractor().extract(content)
            if suffix in {".png", ".jpg", ".jpeg"}:
                return ImageExtractor().extract(content, filename)
            return ExtractionResult(text="", status="failed", error=f"Unsupported type: {suffix}")
        except Exception as exc:  # noqa: BLE001 — surface extractor failures to the record
            return ExtractionResult(text="", status="failed", error=str(exc))
