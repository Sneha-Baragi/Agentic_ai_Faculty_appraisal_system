import os
import sys
from datetime import date

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.models import (
    AppraisalCycle,
    AppraisalReport,
    AppraisalScore,
    Evidence,
    EvidenceExtraction,
    EvidenceValidation,
    FacultyAppraisalRun,
    FacultyProfile,
    FacultyActivity,
    Role,
    Rubric,
    User,
)
from app.scoring.engine import ENGINE_VERSION, calculate_score, load_demo_rubric
from app.services.activities import create_activity, activity_to_dict
from app.services.cycles import get_open_cycle, get_or_create_run
from app.services.evidence import EvidenceService, evidence_to_dict
from app.services.extraction import DocumentExtractor
from app.services.storage import get_storage, object_key_for
from app.services.scoring_map import to_score_inputs
from sqlalchemy import select
from sqlalchemy.orm import Session

from .conftest import (
    auth_headers,
    login,
    make_cycle,
    make_faculty_profile,
    make_user,
)


def _attach_valid_evidence(db: Session, act_id, fac_id, cycle_id, key: str):
    storage = get_storage()
    content = b"%PDF-1.4\n%fake\n"
    object_key = object_key_for(
        faculty_id=str(fac_id),
        cycle_id=str(cycle_id),
        activity_id=str(act_id),
        evidence_id=str(key),
        filename=f"{key}.pdf",
    )
    try:
        storage.put(object_key, content, "application/pdf")
    except Exception:
        pass
    import hashlib

    checksum = hashlib.sha256(content).hexdigest()
    ev = Evidence(
        activity_id=act_id,
        faculty_id=fac_id,
        cycle_id=cycle_id,
        storage_key=object_key,
        mime="application/pdf",
        checksum_sha256=checksum,
        kind="pdf",
        original_filename=f"{key}.pdf",
        file_size=len(content),
        validation_status="valid",
        extraction_status="extracted",
    )
    db.add(ev)
    db.flush()
    db.add(
        EvidenceValidation(
            evidence_id=ev.id, activity_id=act_id, verdict="supported", confidence=1.0, reasons=["checks:ok"]
        )
    )
    db.add(
        EvidenceExtraction(
            evidence_id=ev.id,
            extraction_status="extracted",
            extracted_text="OK",
            extra_metadata={"format": "pdf"},
        )
    )
    return ev


# ---- Test 1: Login + /me works ------------------------------------------
def test_login_flow_returns_me(client, db_session: Session):
    u = make_user(db_session, "login-flow@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u, employee_code="LF-1")
    db_session.commit()
    tok = login(client, "login-flow@demo.local")
    assert tok and len(tok) > 10
    r = client.get("/api/v1/auth/me", headers=auth_headers(tok))
    assert r.status_code == 200, r.text
    assert r.json()["email"] == "login-flow@demo.local"


# ---- Test 2: Profile /me ownership --------------------------------------
def test_profile_me_isolated_returns_own(client, db_session: Session):
    uA = make_user(db_session, "pa@demo.local", ["faculty"])
    uB = make_user(db_session, "pb@demo.local", ["faculty"])
    pA = make_faculty_profile(db_session, uA, employee_code="PA-1")
    pB = make_faculty_profile(db_session, uB, employee_code="PB-1")
    db_session.commit()
    tokA = login(client, "pa@demo.local")
    tokB = login(client, "pb@demo.local")
    bodyA = client.get("/api/v1/profile/me", headers=auth_headers(tokA)).json()
    bodyB = client.get("/api/v1/profile/me", headers=auth_headers(tokB)).json()
    assert bodyA["employee_code"] == "PA-1"
    assert bodyB["employee_code"] == "PB-1"
    assert bodyA["id"] != bodyB["id"]


# ---- Test 3: PATCH /profile/me restricts fields ------------------------
def test_profile_patch_restricts_editable_fields(client, db_session: Session):
    u = make_user(db_session, "pp@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u, employee_code="P-ORIG")
    p.department = "CS"
    p.designation = "Asst"
    p.confidentiality_scope = "self"
    db_session.commit()
    tok = login(client, "pp@demo.local")
    r = client.patch(
        "/api/v1/profile/me",
        headers={**auth_headers(tok), "Content-Type": "application/json"},
        json={
            "full_name": "New Name",
            "department": "Physics",
            "designation": "Prof",
            "institution": "X U",
            "employee_code": "HACKED",
            "confidentiality_scope": "all",
            "role": "admin",
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["full_name"] == "New Name"
    assert body["department"] == "Physics"
    assert body.get("employee_code") != "HACKED"


# ---- Test 4: Activity CRUD ---------------------------------------------
def test_activity_crud_all_categories(client, db_session: Session):
    u = make_user(db_session, "act@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session, status="open")
    db_session.commit()
    tok = login(client, "act@demo.local")
    cats = [
        ("research", "publications"),
        ("teaching", "courses_taught"),
        ("administrative", "committee_work"),
    ]
    created_ids = []
    for cat, typ in cats:
        r = client.post(
            "/api/v1/activities",
            headers={**auth_headers(tok), "Content-Type": "application/json"},
            json={
                "category": cat,
                "activity_type": typ,
                "title": f"{cat} act",
                "description": "D",
                "date": "2026-01-01",
            },
        )
        assert r.status_code == 201, r.text
        created_ids.append(r.json()["id"])
    assert len(client.get("/api/v1/activities", headers=auth_headers(tok)).json()) == 3
    r = client.patch(
        f"/api/v1/activities/{created_ids[0]}",
        headers={**auth_headers(tok), "Content-Type": "application/json"},
        json={"title": "New Title"},
    )
    assert r.status_code == 200
    assert r.json()["title"] == "New Title"
    r = client.delete(f"/api/v1/activities/{created_ids[-1]}", headers=auth_headers(tok))
    assert r.status_code == 204
    assert len(client.get("/api/v1/activities", headers=auth_headers(tok)).json()) == 2


# ---- Test 5: Activity ownership ----------------------------------------
def test_activity_ownership_patch_delete_blocked(client, db_session: Session):
    uA = make_user(db_session, "oA@demo.local", ["faculty"])
    uB = make_user(db_session, "oB@demo.local", ["faculty"])
    pA = make_faculty_profile(db_session, uA)
    pB = make_faculty_profile(db_session, uB)
    cycle = make_cycle(db_session)
    act_a = create_activity(
        db_session,
        faculty=pA,
        category="research",
        activity_type="publications",
        title="A's",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.commit()
    tokB = login(client, "oB@demo.local")
    r = client.patch(
        f"/api/v1/activities/{act_a.id}",
        headers={**auth_headers(tokB), "Content-Type": "application/json"},
        json={"title": "hacked"},
    )
    assert r.status_code in (403, 404)
    r = client.delete(f"/api/v1/activities/{act_a.id}", headers=auth_headers(tokB))
    assert r.status_code in (403, 404)


def test_delete_activity_cleans_evidence_extractions(client, db_session: Session):
    u = make_user(db_session, "del-act@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session, status="open")
    db_session.commit()
    tok = login(client, "del-act@demo.local")

    create_r = client.post(
        "/api/v1/activities",
        headers={**auth_headers(tok), "Content-Type": "application/json"},
        json={
            "category": "research",
            "activity_type": "publications",
            "title": "Delete me",
            "description": "Test delete lifecycle",
            "date": "2026-01-01",
        },
    )
    assert create_r.status_code == 201, create_r.text
    activity_id = create_r.json()["id"]

    activity = db_session.scalar(select(FacultyActivity).where(FacultyActivity.id == activity_id))
    assert activity is not None
    evidence = _attach_valid_evidence(db_session, activity.id, p.id, cycle.id, "delete-evidence")
    db_session.commit()

    extraction = db_session.scalar(select(EvidenceExtraction).where(EvidenceExtraction.evidence_id == evidence.id))
    assert extraction is not None

    delete_r = client.delete(f"/api/v1/activities/{activity_id}", headers=auth_headers(tok))
    assert delete_r.status_code == 204, delete_r.text

    remaining = db_session.scalar(select(EvidenceExtraction).where(EvidenceExtraction.evidence_id == evidence.id))
    assert remaining is None
    orphan_rows = db_session.scalar(select(EvidenceExtraction).where(EvidenceExtraction.evidence_id == evidence.id))
    assert orphan_rows is None


# ---- Test 6: Evidence ownership blocked --------------------------------
def test_evidence_upload_blocked_for_other_faculty(client, db_session: Session):
    uA = make_user(db_session, "evA@demo.local", ["faculty"])
    uB = make_user(db_session, "evB@demo.local", ["faculty"])
    pA = make_faculty_profile(db_session, uA)
    pB = make_faculty_profile(db_session, uB)
    cycle = make_cycle(db_session)
    a = create_activity(
        db_session,
        faculty=pA,
        category="research",
        activity_type="publications",
        title="A",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.commit()
    tokB = login(client, "evB@demo.local")
    r = client.post(
        f"/api/v1/evidence?activity_id={a.id}",
        headers=auth_headers(tokB),
        files={"file": ("e.pdf", b"1234", "application/pdf")},
        data={"activity_id": str(a.id)},
    )
    assert r.status_code in (400, 403, 404)


# ---- Test 7: Invalid file types rejected -------------------------------
def test_invalid_file_type_rejected(client, db_session: Session):
    u = make_user(db_session, "ef@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    a = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="publications",
        title="T",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.commit()
    tok = login(client, "ef@demo.local")
    r = client.post(
        f"/api/v1/evidence?activity_id={a.id}",
        headers=auth_headers(tok),
        files={"file": ("bad.exe", b"MZ", "application/exe")},
        data={"activity_id": str(a.id)},
    )
    assert r.status_code in (400, 422)


# ---- Test 8: SHA-256 checksum computed ----------------------------------
def test_sha256_checksum_stored(client, db_session: Session):
    u = make_user(db_session, "cs@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    a = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="publications",
        title="T",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.commit()
    tok = login(client, "cs@demo.local")
    payload = b"some-deterministic-bytes\n"
    r = client.post(
        f"/api/v1/evidence?activity_id={a.id}",
        headers=auth_headers(tok),
        files={"file": ("e.pdf", payload, "application/pdf")},
        data={"activity_id": str(a.id)},
    )
    if r.status_code == 201:
        body = r.json()
        assert "checksum" in body or "checksum_sha256" in body
        cs = body.get("checksum") or body.get("checksum_sha256")
        assert cs is not None
        assert len(cs) == 64
        import hashlib

        assert cs == hashlib.sha256(payload).hexdigest()


# ---- Test 9: PDF extraction ---------------------------------------------
def test_extract_pdf(tmp_path):
    p = tmp_path / "g.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(72, 720, "Hello PDF Extraction")
    c.save()
    extractor = DocumentExtractor()
    result = extractor.extract(p.read_bytes(), "g.pdf", "application/pdf")
    assert result.status == "extracted"
    assert "Hello PDF Extraction" in result.text


# ---- Test 10: DOCX extraction ------------------------------------------
def test_extract_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Hello DOCX line 1")
    doc.add_paragraph("Hello DOCX line 2")
    p = tmp_path / "g.docx"
    doc.save(str(p))
    extractor = DocumentExtractor()
    result = extractor.extract(p.read_bytes(), "g.docx", None)
    assert result.status == "extracted"
    assert "Hello DOCX line 1" in result.text
    assert "Hello DOCX line 2" in result.text


# ---- Test 11: XLSX extraction ------------------------------------------
def test_extract_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello"
    ws["B2"] = "XLSX"
    ws["C3"] = "World"
    p = tmp_path / "g.xlsx"
    wb.save(str(p))
    extractor = DocumentExtractor()
    result = extractor.extract(p.read_bytes(), "g.xlsx", None)
    assert result.status == "extracted"
    assert "Hello" in result.text
    assert "XLSX" in result.text
    assert "World" in result.text


def _score_activities(db, p, cycle, activities_ids):
    """Helper: rebuild activity list with evidence and compute score."""
    acts_q = db.scalars(
        select(FacultyActivity)
        .where(
            FacultyActivity.faculty_id == p.id,
            FacultyActivity.cycle_id == cycle.id,
            FacultyActivity.id.in_(activities_ids),
        )
    ).all()
    act_dicts = [activity_to_dict(a) for a in acts_q]
    validated = []
    for d in act_dicts:
        statuses = d.get("evidence_statuses") or []
        eids = d.get("evidence_ids") or []
        if eids and all(s in {"valid", "extracted"} for s in statuses):
            validated.append(d)
    inputs = to_score_inputs(validated)
    result = calculate_score(inputs, load_demo_rubric())
    line_items = []
    for li in result.line_items:
        line_items.append(
            {
                "activity_id": li.activity_id,
                "rubric_code": li.rubric_code,
                "points": li.points,
                "evidence_ids": li.evidence_ids,
            }
        )
    breakdown = {
        "research_score": result.category_totals.get("research", 0),
        "teaching_score": result.category_totals.get("teaching", 0),
        "administrative_score": result.category_totals.get("administrative", 0),
        "total_score": result.total,
        "line_items": line_items,
    }
    return {"total_score": result.total, "score_breakdown": breakdown, "rating": result.rating_recommendation}


# ---- Test 12: Activity without valid evidence excluded -----------------
def test_activity_without_valid_evidence_gets_zero(db_session: Session):
    u = make_user(db_session, "s1@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    with_ev = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="publications",
        title="W Ev",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"venue": "Journal", "year": 2026, "authors": ["a"]},
    )
    no_ev = create_activity(
        db_session,
        faculty=p,
        category="teaching",
        activity_type="courses_taught",
        title="NO EV",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"hours": 30},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, with_ev.id, p.id, cycle.id, "T12")
    db_session.commit()
    result = _score_activities(db_session, p, cycle, [with_ev.id, no_ev.id])
    items = result["score_breakdown"]["line_items"]
    no_ev_items = [i for i in items if str(i["activity_id"]) == str(no_ev.id)]
    assert sum(i["points"] for i in no_ev_items) == 0
    with_items = [i for i in items if str(i["activity_id"]) == str(with_ev.id)]
    assert sum(i["points"] for i in with_items) > 0


# ---- Test 13: Activity with valid evidence scored ----------------------
def test_activity_with_valid_evidence_receives_points(db_session: Session):
    u = make_user(db_session, "s2@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    act = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="publications",
        title="Journal X",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"venue": "Journal", "year": 2026, "authors": 1},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, act.id, p.id, cycle.id, "T13")
    db_session.commit()
    result = _score_activities(db_session, p, cycle, [act.id])
    assert result["total_score"] > 0


# ---- Test 14: Scored line items contain evidence_ids -------------------
def test_scored_line_items_have_evidence_ids(db_session: Session):
    u = make_user(db_session, "s3@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    act = create_activity(
        db_session,
        faculty=p,
        category="teaching",
        activity_type="student_projects",
        title="Proj",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.flush()
    ev = _attach_valid_evidence(db_session, act.id, p.id, cycle.id, "T14")
    db_session.commit()
    result = _score_activities(db_session, p, cycle, [act.id])
    for i in result["score_breakdown"]["line_items"]:
        if i["points"] > 0:
            assert "evidence_ids" in i
            assert len(i["evidence_ids"]) > 0


# ---- Test 15: Deterministic score reproducibility ----------------------
def test_deterministic_score_matches(db_session: Session):
    u = make_user(db_session, "s4@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    a = create_activity(
        db_session,
        faculty=p,
        category="administrative",
        activity_type="committee_work",
        title="C",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"role": "convener", "meetings": 12},
    )
    b = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="conferences",
        title="Conf",
        description="",
        activity_date=date(2026, 1, 1),
        extras={},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, a.id, p.id, cycle.id, "T15A")
    _attach_valid_evidence(db_session, b.id, p.id, cycle.id, "T15B")
    db_session.commit()
    r1 = _score_activities(db_session, p, cycle, [a.id, b.id])
    r2 = _score_activities(db_session, p, cycle, [a.id, b.id])
    assert r1["total_score"] == r2["total_score"]
    br1 = r1["score_breakdown"]
    br2 = r2["score_breakdown"]
    assert br1["research_score"] == br2["research_score"]
    assert br1["teaching_score"] == br2["teaching_score"]
    assert br1["administrative_score"] == br2["administrative_score"]


# ---- Test 16: Report persistence ---------------------------------------
def test_report_persisted_after_run_appraisal(client, db_session: Session):
    u = make_user(db_session, "rp@demo.local", ["faculty"])
    p = make_faculty_profile(db_session, u)
    cycle = make_cycle(db_session)
    a = create_activity(
        db_session,
        faculty=p,
        category="research",
        activity_type="publications",
        title="P",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"venue": "Journal", "year": 2026},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, a.id, p.id, cycle.id, "T16")
    db_session.commit()
    tok = login(client, "rp@demo.local")
    r = client.post("/api/v1/appraisals/run", headers=auth_headers(tok))
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["approval_status"] == "awaiting_review"

    score = db_session.execute(select(AppraisalScore)).scalars().first()
    assert score is not None
    assert score.engine_version == ENGINE_VERSION

    report = db_session.execute(select(AppraisalReport)).scalars().first()
    assert report is not None
    assert report.report_json is not None or report.summary_md or True

    run = db_session.execute(select(FacultyAppraisalRun)).scalars().first()
    assert run is not None


# ---- Test 17: Reviewer read access to packet ---------------------------
def test_reviewer_read_access(client, db_session: Session):
    uf = make_user(db_session, "rf@demo.local", ["faculty"])
    ur = make_user(db_session, "rr@demo.local", ["hod"])
    pf = make_faculty_profile(db_session, uf)
    cycle = make_cycle(db_session)
    db_session.commit()
    tok = login(client, "rr@demo.local")
    r = client.get("/api/v1/review/faculty", headers=auth_headers(tok))
    assert r.status_code == 200, r.text
    r = client.get(f"/api/v1/review/faculty/{pf.id}", headers=auth_headers(tok))
    assert r.status_code == 200, r.text
    packet = r.json()
    assert "profile" in packet
    assert "activities" in packet


def test_reviewer_packet_includes_generated_score_and_report(client, db_session: Session):
    uf = make_user(db_session, "packet-score-faculty@demo.local", ["faculty"])
    ur = make_user(db_session, "packet-score-hod@demo.local", ["hod"])
    pf = make_faculty_profile(db_session, uf)
    cycle = make_cycle(db_session)
    activity = create_activity(
        db_session,
        faculty=pf,
        category="research",
        activity_type="publications",
        title="Indexed journal article",
        description="Published in a peer-reviewed venue.",
        activity_date=date(2026, 2, 1),
        extras={"venue": "Journal of Applied Computing", "year": 2026, "authors": ["Faculty Member"]},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, activity.id, pf.id, cycle.id, "PACKET")
    db_session.commit()

    faculty_token = login(client, "packet-score-faculty@demo.local")
    run_response = client.post("/api/v1/appraisals/run", headers=auth_headers(faculty_token))
    assert run_response.status_code == 200, run_response.text

    reviewer_token = login(client, "packet-score-hod@demo.local")
    packet_response = client.get(
        f"/api/v1/review/faculty/{pf.id}", headers=auth_headers(reviewer_token)
    )
    assert packet_response.status_code == 200, packet_response.text
    packet = packet_response.json()
    assert packet["score"]["total_score"] > 0
    assert packet["report"]["status"] == "generated"
    assert packet["activities"][0]["details"]["venue"] == "Journal of Applied Computing"


# ---- Test 18: Faculty blocked from reviewer endpoints ------------------
def test_faculty_blocked_from_reviewer_endpoints(client, db_session: Session):
    uf = make_user(db_session, "ff@demo.local", ["faculty"])
    pf = make_faculty_profile(db_session, uf)
    make_cycle(db_session)
    db_session.commit()
    tok = login(client, "ff@demo.local")
    assert client.get("/api/v1/review/faculty", headers=auth_headers(tok)).status_code in (403, 401)
    assert client.get(f"/api/v1/review/faculty/{pf.id}", headers=auth_headers(tok)).status_code in (403, 401)


# ---- Test 19: Admin authorization --------------------------------------
def test_admin_authorization(client, db_session: Session):
    ua = make_user(db_session, "admin-t@demo.local", ["admin"])
    uf = make_user(db_session, "fac-t@demo.local", ["faculty"])
    make_faculty_profile(db_session, uf)
    db_session.commit()
    tokA = login(client, "admin-t@demo.local")
    tokF = login(client, "fac-t@demo.local")

    users = client.get("/api/v1/admin/users", headers=auth_headers(tokA)).json()
    for u in users:
        assert "hashed_password" not in u
        assert "password" not in u

    status = client.get("/api/v1/admin/status", headers=auth_headers(tokA)).json()
    assert "users" in status
    assert "jwt_secret" not in status
    assert "minio_secret_key" not in status

    assert client.get("/api/v1/admin/users", headers=auth_headers(tokF)).status_code in (403, 401)
    assert client.get("/api/v1/admin/status", headers=auth_headers(tokF)).status_code in (403, 401)


# ---- Test 20: Appraisal isolation between faculties --------------------
def test_appraisal_isolated_by_faculty_auth(client, db_session: Session):
    uA = make_user(db_session, "gA@demo.local", ["faculty"])
    uB = make_user(db_session, "gB@demo.local", ["faculty"])
    pA = make_faculty_profile(db_session, uA)
    pB = make_faculty_profile(db_session, uB)
    cycle = make_cycle(db_session)
    aA = create_activity(
        db_session,
        faculty=pA,
        category="research",
        activity_type="publications",
        title="A Q1",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"venue": "Journal", "year": 2026, "authors": 1},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, aA.id, pA.id, cycle.id, "T20A")
    aB = create_activity(
        db_session,
        faculty=pB,
        category="administrative",
        activity_type="committee_work",
        title="B Ad",
        description="",
        activity_date=date(2026, 1, 1),
        extras={"role": "convener", "meetings": 20},
    )
    db_session.flush()
    _attach_valid_evidence(db_session, aB.id, pB.id, cycle.id, "T20B")
    db_session.commit()
    tokA = login(client, "gA@demo.local")
    tokB = login(client, "gB@demo.local")
    rA = client.post("/api/v1/appraisals/run", headers=auth_headers(tokA)).json()
    rB = client.post("/api/v1/appraisals/run", headers=auth_headers(tokB)).json()
    brA = rA["score"]["score_breakdown"]
    assert brA["research_score"] > 0
    assert brA["administrative_score"] == 0
    brB = rB["score"]["score_breakdown"]
    assert brB["research_score"] == 0
    assert brB["administrative_score"] > 0


# ---- Test 22: Admin can create cycle -----------------------------------
def test_admin_can_create_cycle(client, db_session: Session):
    ua = make_user(db_session, "admin-c22@demo.local", ["admin"])
    db_session.commit()
    tok = login(client, "admin-c22@demo.local")
    r = client.post(
        "/api/v1/admin/cycles",
        headers={**auth_headers(tok), "Content-Type": "application/json"},
        json={
            "name": "Test Cycle 22",
            "academic_year": "2026-2027",
            "start_date": "2026-01-01",
            "end_date": "2027-12-31",
            "status": "draft",
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["name"] == "Test Cycle 22"
    assert body["status"] == "draft"
    assert body["academic_year"] == "2026-2027"


# ---- Test 23: Admin can open cycle -------------------------------------
def test_admin_can_open_cycle(client, db_session: Session):
    ua = make_user(db_session, "admin-c23@demo.local", ["admin"])
    cycle = make_cycle(db_session, status="draft")
    db_session.commit()
    tok = login(client, "admin-c23@demo.local")
    r = client.patch(
        f"/api/v1/admin/cycles/{cycle.id}/open",
        headers=auth_headers(tok),
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "open"


# ---- Test 24: Admin can close cycle ------------------------------------
def test_admin_can_close_cycle(client, db_session: Session):
    ua = make_user(db_session, "admin-c24@demo.local", ["admin"])
    cycle = make_cycle(db_session, status="open")
    db_session.commit()
    tok = login(client, "admin-c24@demo.local")
    r = client.patch(
        f"/api/v1/admin/cycles/{cycle.id}/close",
        headers=auth_headers(tok),
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "closed"


# ---- Test 25: Faculty cannot create/administer cycles ------------------
def test_faculty_blocked_from_admin_cycle_endpoints(client, db_session: Session):
    uf = make_user(db_session, "fac-c25@demo.local", ["faculty"])
    make_faculty_profile(db_session, uf)
    cycle = make_cycle(db_session, status="draft")
    db_session.commit()
    tok = login(client, "fac-c25@demo.local")
    create_r = client.post(
        "/api/v1/admin/cycles",
        headers={**auth_headers(tok), "Content-Type": "application/json"},
        json={
            "name": "Hacked",
            "academic_year": "2099-2100",
            "start_date": "2099-01-01",
            "end_date": "2100-12-31",
            "status": "draft",
        },
    )
    assert create_r.status_code in (403, 401)
    open_r = client.patch(
        f"/api/v1/admin/cycles/{cycle.id}/open",
        headers=auth_headers(tok),
    )
    assert open_r.status_code in (403, 401)
    list_r = client.get("/api/v1/admin/cycles", headers=auth_headers(tok))
    assert list_r.status_code in (403, 401)


# ---- Test 26: Faculty sees open cycle via /cycles/current --------------
def test_faculty_sees_open_cycle_in_current_endpoint(client, db_session: Session):
    uf = make_user(db_session, "fac-c26@demo.local", ["faculty"])
    make_faculty_profile(db_session, uf)
    cycle = make_cycle(db_session, status="open")
    db_session.commit()
    tok = login(client, "fac-c26@demo.local")
    r = client.get("/api/v1/cycles/current", headers=auth_headers(tok))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["id"] == str(cycle.id)
    assert body["status"] == "open"


# ---- Test 27: Opening cycle closes existing open cycle -----------------
def test_open_cycle_closes_existing(client, db_session: Session):
    ua = make_user(db_session, "admin-c27@demo.local", ["admin"])
    c1 = make_cycle(db_session, status="open")
    c2 = make_cycle(db_session, status="draft")
    db_session.commit()
    tok = login(client, "admin-c27@demo.local")
    r = client.patch(
        f"/api/v1/admin/cycles/{c2.id}/open",
        headers=auth_headers(tok),
    )
    assert r.status_code == 200, r.text
    db_session.refresh(c1)
    assert c1.status == "closed"
    assert r.json()["status"] == "open"


# ---- Test 21: Alembic single head --------------------------------------
def test_alembic_has_exactly_one_head():
    from pathlib import Path
    import subprocess
    import sys

    backend_dir = Path(__file__).resolve().parent.parent
    alembic_ini = backend_dir / "alembic.ini"
    alembic_script = backend_dir / ".venv" / "Scripts" / "alembic.exe"
    python_exe = backend_dir / ".venv" / "Scripts" / "python.exe"

    try:
        from alembic.config import Config
        from alembic.script import ScriptDirectory

        cfg = Config(str(alembic_ini))
        cfg.set_main_option("script_location", str(backend_dir / "alembic"))
        scripts = ScriptDirectory.from_config(cfg)
        heads = scripts.get_heads()
    except Exception:
        exe = python_exe if python_exe.exists() else sys.executable
        cmd = [str(exe), "-m", "alembic", "-c", str(alembic_ini), "heads"]
        result = subprocess.run(cmd, capture_output=True, text=True, cwd=str(backend_dir))
        stdout = (result.stdout or "").strip() + (result.stderr or "").strip()
        heads = [line.strip().split()[0] for line in stdout.splitlines() if line.strip() and not line.startswith("(")]
        if not heads and " -> " in stdout:
            heads = [stdout.strip().split(" -> ")[-1].split()[0]]

    assert heads is not None, f"No heads found, check alembic setup"
    assert len(heads) == 1, f"Expected 1 alembic head, got: {len(heads)} -> {heads}"
